"""
Lectura de las planillas oficiales del ranking (FEJUCHILE).

Cada archivo "RANKING <GRUPO> <AÑO>.xlsx" trae una hoja por categoría de peso.
Formato verificado con las planillas 2026:

  * En la fila de encabezados ("APELLIDOS") cada año tiene un bloque de columnas
    que empieza con una celda fecha 1-ene-AAAA y termina en una columna "TOTAL".
  * La primera columna del bloque es el ARRASTRE: 50 % de los puntos ganados en
    torneos del año anterior. Luego vienen los torneos del año.
  * TOTAL del año = arrastre + puntos de torneos. Es lo que ordena el ranking.
  * Absoluta, Sub 21 y Sub 18 traen un segundo bloque del mismo año con el
    ranking internacional, que funciona igual.
"""
from __future__ import annotations

import datetime as dt
import logging
import os
import re
import unicodedata
from pathlib import Path

import openpyxl
import pandas as pd

from . import ubicaciones

log = logging.getLogger(__name__)

# clave en el nombre del archivo -> (etiqueta, orden, slug)
GRUPOS = {
    "SUB13": ("Sub 13", 1, "sub-13"),
    "SUB15": ("Sub 15", 2, "sub-15"),
    "SUB18": ("Sub 18", 3, "sub-18"),
    "SUB21": ("Sub 21", 4, "sub-21"),
    "ABSOLUTA": ("Absoluta", 5, "absoluta"),
}
# Edades de menores de 18 años (Sub 13, Sub 15 y Sub 18; se calcula sola si se agrega otra edad).
# Sus fichas y rankings no se indexan en buscadores, y en las demás páginas Google no muestra sus nombres.
GRUPOS_MENORES = {slug for etiqueta, _, slug in GRUPOS.values()
                  if etiqueta.startswith("Sub ") and int(etiqueta[4:]) <= 18}
SEXOS = {"F": "Mujeres", "M": "Hombres"}

# Variantes de escritura del mismo club detectadas en las planillas.
# Si aparece otra, agrégala aquí: "COMO APARECE": "NOMBRE UNIFICADO".
ALIAS_CLUBES = {
    "ACADEMIA ROMO": "ACAD ROMO",
    "ARUKIKATA": "ARUKI KATA", "AUKI KATA": "ARUKI KATA",
    "CARB ARICA": "CARAB ARICA",
    "CARABINEROS CHILE": "CARABINEROS", "CARABINEROS DE CHILE": "CARABINEROS",
    "CARBINEROS DE CHILE": "CARABINEROS",
    "CCSH": "CCRSH", "CRSH": "CCRSH",
    "CHOLCHOL": "CHOL CHOL",
    "COL REP SIRIA": "COLEGIO REP SIRIA", "REP SIRIA": "COLEGIO REP SIRIA",
    "E I CORDOVA": "ESC ITALO CORDOVA", "ESC I CORDOVA": "ESC ITALO CORDOVA",
    "ESCUELA ITALO CORDOVA": "ESC ITALO CORDOVA",
    "ESC SUB OF EJERCITO": "ESC SUBOFICIALES EJERCITO",
    "ESC SUB OFICIALES": "ESC SUBOFICIALES EJERCITO",
    "ESC-SUBOF EJERCITO": "ESC SUBOFICIALES EJERCITO",
    "ESSOF": "ESC SUBOFICIALES EJERCITO",
    "ESCPDI": "ESCUELA PDI",
    "ESLABÓN": "ESLABON",
    "ESTAMAN": "EASTMAN",
    "FUNDACIO AEDAR": "FUNDACION AEDAR",
    "GUZMAN": "JUDO GUZMAN",
    "ISC SAN BERNADO": "ISC SAN BERNARDO",
    "JIAME CHARLES": "JAIME CHARLES",
    "JUDO QUILICRUA": "JUDO QUILICURA", "QUILICURA": "JUDO QUILICURA",
    "JUDO TRUNO": "JUDO TRUENO", "JUDI TRUENO": "JUDO TRUENO", "TRUENO": "JUDO TRUENO",
    "JUDO NUBLE": "JUDO ÑUBLE", "ÑUBLE": "JUDO ÑUBLE",
    "JUDO CARBON": "JUDO DEL CARBON", "EL CARBON": "JUDO DEL CARBON",
    "JUDO TALCA": "TALCA",
    "KAZOKU LA SERENA": "KAZOKU LS",
    "KIDO": "KIDO CURACAVI", "KIDOC CURACAVI": "KIDO CURACAVI",
    "KITSUGI": "KINTSUGI",
    "KODOKAN": "KODOKAN CHILE",
    "LEONES DE TARAPACA": "LEONES TARAPACA",
    "MANKE PUCON": "MANKE",
    "OLYMPIC DREAM": "OLYMPIC DREAMS",
    "PINTANA": "LA PINTANA",
    "PGN": "PROY G NAVIA", "PROY. G NAVIA": "PROY G NAVIA",
    "RISEI AKURA": "RISEI SAKURA",
    # "SAN SEI" (La Reina) y "SAN SEI PTO VARAS" (Puerto Varas) son clubes distintos en el registro de la federación.
    "SAMN SEI": "SAN SEI", "SANSEI": "SAN SEI",
    "SAIGO KIDS": "SHIRO SAIGO KIDS",
    "SHIN GI TAI PTO NATALES": "SHIN GI TAI",
    "SHOGAI": "SHOGAI ÑUÑOA",
    "UNIV ANTOFAGASTA": "U ANTOFAGASTA",
    "U CATOLICA DE NORTE": "U CATOLICA DEL NORTE", "U. CATOLICA DEL NORTE": "U CATOLICA DEL NORTE",
    "U CHILE": "U DE CHILE", "UNIVERSIDAD DE CHILE": "U DE CHILE",
    "VIÑA": "VIÑA DEL MAR",
    "YAMAARASHI ANCUD": "YAMA ARASHI ANCUD",
    "YAMARASHI": "YAMA HARASHI",
    # Confirmadas porque un mismo deportista aparece con ambos nombres, o por el registro de clubes:
    "CHEISUKAN": "CHUSEIKAN",
    "TOME": "KAZOKU TOME",
    "KAZOKU JC": "KAZOKU LS",
    "SHIJOUKI": "SHOUJIKI",
    "CASABLANCA": "JITA KYOEI",
    "PUC": "UC",
    # Clubes que cambiaron de nombre o que la federación escribe con siglas (registro de clubes de FEJUCHILE):
    "OZANAM": "CHUSEIKAN",           # Club de Judo Ozanam pasó a llamarse Chuseikan
    "ALEMAN": "CDUSS",               # hoy Club Deportivo Universidad San Sebastián
    "EIM": "ESTADIO ISRAELITA",      # Estadio Israelita Maccabi
    "AUSTRAL": "UACH",               # Universidad Austral de Chile
    "PDI": "ESCUELA PDI",
}
# Filas sin club en la planilla. ("NAC" sí es un club: el North American College de Arica.)
SIN_CLUB = {""}
SLUG_SIN_CLUB = "sin-club"


class PlanillaInvalida(ValueError):
    """La planilla no tiene el formato esperado."""


# ----------------------------------------------------------------------------
# Utilidades de texto
# ----------------------------------------------------------------------------
def limpiar(txt) -> str:
    """Texto de una celda en mayúsculas y sin espacios de más (celda vacía -> '')."""
    if txt is None:
        return ""
    return re.sub(r"\s+", " ", str(txt)).strip().upper()


def sin_tildes(txt: str) -> str:
    """'Ñuñoa Peñalolén' -> 'Nunoa Penalolen'."""
    t = unicodedata.normalize("NFD", txt)
    return "".join(ch for ch in t if unicodedata.category(ch) != "Mn")


def slug(txt: str) -> str:
    """Texto para direcciones web: 'Río Cruces' -> 'rio-cruces'."""
    return re.sub(r"[^a-z0-9]+", "-", sin_tildes(txt.lower())).strip("-")


def normalizar_club(txt) -> str:
    """Nombre unificado del club: sin puntos ni espacios de más y con las variantes de ALIAS_CLUBES resueltas."""
    c = re.sub(r"\s+", " ", limpiar(txt).replace(".", "")).strip()
    return ALIAS_CLUBES.get(c, ALIAS_CLUBES.get(limpiar(txt), c))


def nombre_club(c: str) -> str:
    """Nombre para mostrar de un club sin ficha en clubes.csv: 'RIO CRUCES' -> 'Rio Cruces'."""
    return "Sin club" if c in SIN_CLUB else c.title()


def slug_club(c: str) -> str:
    """Identificador del club en las direcciones web."""
    return SLUG_SIN_CLUB if c in SIN_CLUB else slug(c)


def detectar_grupo(nombre_archivo: str) -> str | None:
    """Grupo de edad según el nombre del archivo: 'RANKING  SUB13 2026.xlsx' -> 'SUB13' (None si no es una planilla)."""
    clave = re.sub(r"[^A-Z0-9]", "", sin_tildes(nombre_archivo.upper()))
    return next((g for g in GRUPOS if g in clave), None)


def _num(v) -> float:
    """Puntaje de una celda: lo que no es número (vacío, texto) cuenta como 0."""
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def _peso(nombre_hoja: str) -> tuple[str, int, str]:
    """'abs+100' -> ('+100 kg', 1100, 'mas100'); '-48' -> ('−48 kg', 48, '48')."""
    p = nombre_hoja.lower().replace("abs", "").replace(" ", "").rstrip("m")
    if not p.startswith(("+", "-")):
        p = "-" + p
    digitos = re.sub(r"\D", "", p)
    if not digitos:
        raise PlanillaInvalida(f"No se reconoce el peso de la hoja '{nombre_hoja}'")
    mas = p.startswith("+")
    return p.replace("-", "−") + " kg", int(digitos) + (1000 if mas else 0), ("mas" if mas else "") + digitos


# ----------------------------------------------------------------------------
# Lectura de hojas
# ----------------------------------------------------------------------------
def _bloque(hdr, anio: int, desde: int = 0):
    """(columna_arrastre, columna_total) del bloque de ese año, buscando desde 'desde'."""
    for j in range(desde, len(hdr)):
        v = hdr[j]
        if isinstance(v, dt.datetime) and (v.year, v.month, v.day) == (anio, 1, 1):
            for k in range(j + 1, len(hdr)):
                if isinstance(hdr[k], str) and hdr[k].strip() == "TOTAL":
                    return j, k
    return None


def _anio_hoja(hdr) -> int:
    """El año más reciente con bloque de puntos. El título de la hoja no sirve:
    a veces queda con el año anterior al copiar hojas."""
    anios = [v.year for v in hdr if isinstance(v, dt.datetime) and (v.month, v.day) == (1, 1)]
    if not anios:
        raise PlanillaInvalida("No se encontró el año del ranking")
    return max(anios)


def leer_hoja(ws, grupo: str) -> tuple[int, list[dict]]:
    """Lee una hoja (una categoría de peso). Devuelve (año del ranking, una fila por deportista).

    Arriba de la fila "APELLIDOS" están los datos de la categoría (CATEG., SEXO, PESO: rótulo en la
    columna B y valor en la C). Bajo esa fila vienen dos filas más de encabezado y luego los deportistas:
    A = lugar, B = primer apellido, C = segundo apellido (casi siempre vacío), D = nombre, E = club, y
    después los bloques de puntos de cada año (ver el comienzo del archivo).
    """
    filas = list(ws.iter_rows(values_only=True))
    try:
        h = next(i for i, r in enumerate(filas) if "APELLIDOS" in [limpiar(x) for x in r])
    except StopIteration:
        raise PlanillaInvalida(f"La hoja '{ws.title}' no tiene la fila 'APELLIDOS'") from None
    meta = {limpiar(r[1]): r[2] for r in filas[:h] if r and len(r) > 2 and r[1]}
    hdr = filas[h]
    anio = _anio_hoja(hdr)

    nac = _bloque(hdr, anio)
    if not nac:
        raise PlanillaInvalida(f"La hoja '{ws.title}' no tiene bloque de puntos {anio}")
    previo = _bloque(hdr, anio - 1)
    intl = _bloque(hdr, anio, nac[1] + 1)  # segundo bloque del año = internacional

    etiqueta, orden, gslug = GRUPOS[grupo]
    sexo = limpiar(meta.get("SEXO"))
    if sexo not in SEXOS:
        raise PlanillaInvalida(f"La hoja '{ws.title}' no indica sexo F o M")
    peso, kg, pslug = _peso(ws.title)

    registros = []
    for r in filas[h + 3:]:  # deportistas: tras la fila "APELLIDOS" y sus dos filas de encabezado
        ape1, ape2, nombre = limpiar(r[1]), limpiar(r[2]), limpiar(r[3])
        if not (ape1 or nombre):
            continue
        a, t = nac
        torneos = [_num(r[j]) for j in range(a + 1, t)]
        rec = {
            "grupo": etiqueta, "orden_grupo": orden, "grupo_slug": gslug,
            "sexo": sexo, "peso": peso, "kg": kg,
            "cat_slug": f"{SEXOS[sexo].lower()}-{pslug}",
            "posicion_planilla": int(r[0]) if isinstance(r[0], (int, float)) else None,
            "deportista": " ".join(x for x in (nombre, ape1, ape2) if x).title(),
            "club_original": limpiar(r[4]),
            "club": normalizar_club(r[4]),
            "total": round(_num(r[t]), 2),
            "arrastre": round(_num(r[a]), 2),
            "torneos": round(sum(torneos), 2),
            "n_torneos": sum(1 for x in torneos if x > 0),
            "total_anterior": round(_num(r[previo[1]]), 2) if previo else 0.0,
            "intl": 0.0, "n_intl": 0,
        }
        if intl:
            ia, it = intl
            ev = [_num(r[j]) for j in range(ia + 1, it)]
            rec["intl"] = round(_num(r[it]), 2)
            rec["n_intl"] = sum(1 for x in ev if x > 0)
        registros.append(rec)
    return anio, registros


def cargar_directorio(directorio: str | Path) -> tuple[pd.DataFrame, int]:
    """Lee todas las planillas de ranking del directorio. Devuelve (datos, año)."""
    registros, anios = [], set()
    archivos = sorted(Path(directorio).glob("*.xlsx"))
    for ruta in archivos:
        grupo = detectar_grupo(ruta.name)
        if not grupo:
            log.warning("Se ignora %s: no corresponde a un grupo de edad conocido", ruta.name)
            continue
        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        try:
            for ws in wb.worksheets:
                anio, regs = leer_hoja(ws, grupo)
                anios.add(anio)
                registros.extend(regs)
        except PlanillaInvalida as exc:
            raise PlanillaInvalida(f"{ruta.name}: {exc}") from None
        finally:
            wb.close()
    if not registros:
        raise PlanillaInvalida(f"No hay planillas de ranking en {directorio}")
    anio = max(anios)
    if len(anios) > 1:
        log.warning("Las planillas mezclan años %s; se usa %s", sorted(anios), anio)

    df = pd.DataFrame(registros)
    df["sexo_txt"] = df.sexo.map(SEXOS)
    df["categoria"] = df.grupo + " · " + df.sexo_txt + " · " + df.peso
    df["cat_id"] = df.grupo_slug + "/" + df.cat_slug
    # El lugar se recalcula por puntaje: algunas planillas traen la numeración desordenada.
    df = df.sort_values(["orden_grupo", "sexo", "kg", "total"], ascending=[True, True, True, False], kind="stable")
    df["pos"] = df.groupby("cat_id").cumcount() + 1
    df["club_slug"] = df.club.map(slug_club)
    # Las planillas no traen RUT ni fecha de nacimiento, y casi siempre un solo apellido: un mismo nombre
    # puede ser de dos personas (un adulto y un niño del Sub 13, por ejemplo). Si un nombre aparece en más
    # de un club, cada club tiene su propia ficha, y la ficha avisa que existe la otra.
    df["dep_nombre_slug"] = df.deportista.map(slug) + "-" + df.sexo.str.lower()
    varios_clubes = df.groupby("dep_nombre_slug").club_slug.transform("nunique") > 1
    df["dep_slug"] = df.dep_nombre_slug.where(~varios_clubes, df.dep_nombre_slug + "-" + df.club_slug)
    df["con_club"] = ~df.club.isin(SIN_CLUB)
    # Nombre visible, región y zona salen de la ficha de cada club (judo/clubes.csv).
    fichas = ubicaciones.cargar()
    df["club_nombre"] = df.club.map(lambda c: fichas.get(c, {}).get("nombre") or nombre_club(c))
    df["region"] = df.club.map(lambda c: fichas.get(c, {}).get("region", ""))
    df["zona"] = df.region.map(ubicaciones.zona_de)
    return df.reset_index(drop=True), anio


def fecha_guardado(ruta: str | Path) -> dt.datetime:
    """Cuándo la federación guardó la planilla (dato interno del .xlsx; si falta, la fecha del archivo)."""
    try:
        wb = openpyxl.load_workbook(ruta, read_only=True)
        guardado = wb.properties.modified
        wb.close()
        if guardado:
            return guardado.replace(tzinfo=dt.timezone.utc)
    except Exception:
        pass
    return dt.datetime.fromtimestamp(os.path.getmtime(ruta), dt.timezone.utc)
